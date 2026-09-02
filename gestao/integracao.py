#!/usr/bin/env python3
# =====================================================================
# SISTEMA DE GESTAO CIO CERRADO
# integracao.py  ·  Sympla · Autentique · Resend
#
# Roda fora do banco, no Agendador de Tarefas (mesmo lugar do
# rotina_cerrado.py). Fala com o Supabase pela service_role, entao
# ignora RLS de proposito: e um processo de servidor, nao um usuario.
#
# Uso:
#   python integracao.py --tudo
#   python integracao.py --sympla          sincroniza inscricoes
#   python integracao.py --jantares        convidados de jantar (link do Sympla de cada jantar)
#   python integracao.py --contratos       envia contratos pendentes
#   python integracao.py --status          le status no Autentique
#   python integracao.py --lembretes       cobra quem nao assinou
#   python integracao.py --emails          processa a fila de envio
#
# Por padrao roda em modo seguro (nao envia nada de verdade).
# Use --producao para valer.
#
# --contratos precisa, alem de SYMPLA_TOKEN/AUTENTIQUE_TOKEN:
#   CONTRATO_TEMPLATE_PATH  caminho do .docx do contrato (com os
#                           placeholders "Prezado(a) participante",
#                           "Nome: ____", "CPF: ____" — mesmo formato
#                           do pipeline antigo, EXPERIENCE 2026)
#   pip install python-docx
#   LibreOffice instalado (soffice no PATH) OU Word + `pip install docx2pdf`,
#   pra converter o docx preenchido em PDF antes de subir pro Autentique
# =====================================================================

import os
import re
import sys
import json
import argparse
import logging
import unicodedata
from datetime import datetime, timezone, timedelta

# .env na mesma pasta, se existir — carrega ANTES de ler as variaveis
# abaixo. So preenche o que ja nao estiver no ambiente (override=False),
# entao uma variavel exportada de verdade (Agendador de Tarefas, CI)
# sempre vence o .env local.
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))
except ImportError:
    pass  # sem python-dotenv, so funciona com variaveis ja exportadas

import requests

# ---------------------------------------------------------------------
# CONFIGURACAO  (variaveis de ambiente, nunca no codigo)
# ---------------------------------------------------------------------
SUPABASE_URL      = os.environ.get("SUPABASE_URL", "")
SUPABASE_SERVICE  = os.environ.get("SUPABASE_SERVICE_KEY", "")
SYMPLA_TOKEN      = os.environ.get("SYMPLA_TOKEN", "")
AUTENTIQUE_TOKEN  = os.environ.get("AUTENTIQUE_TOKEN", "")
RESEND_KEY        = os.environ.get("RESEND_API_KEY", "")

EVENTO_SLUG       = os.environ.get("CERRADO_EVENTO", "cerrado2027")
REMETENTE         = os.environ.get("CERRADO_REMETENTE", "contato@ciocerrado.com.br")

# Caminho do .docx do contrato (o mesmo padrao do pipeline antigo,
# cerrado_contratos.py: preenche nome/CPF no docx, converte pra PDF e
# sobe pro Autentique — NAO usa createDocumentFromTemplate, porque
# nunca existiu template cadastrado no Autentique, so o docx local).
# Sem isso configurado, envio de contrato fica pulado (ver enviar_contratos).
CONTRATO_TEMPLATE_PATH = os.environ.get("CONTRATO_TEMPLATE_PATH", "")

# Prazo de assinatura no Autentique (bloqueia assinatura apos esta data).
# Formato: "AAAA-MM-DDTHH:MM:SS.000-03:00". Vazio = sem prazo.
AUTENTIQUE_DEADLINE = os.environ.get("AUTENTIQUE_DEADLINE", "")

# Lembrete de contrato: dias sem assinar antes da primeira cobranca,
# e intervalo minimo entre uma cobranca e a proxima.
LEMBRETE_APOS_DIAS   = 3
LEMBRETE_INTERVALO   = 4
LEMBRETE_MAXIMO      = 3

SCHEMA = "gestao"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-7s %(message)s",
    datefmt="%d/%m %H:%M:%S")
log = logging.getLogger("cerrado")


# =====================================================================
# CAMADA SUPABASE
# =====================================================================
class Supa:
    """Acesso ao PostgREST. Todas as chamadas usam o schema gestao."""

    def __init__(self, url, key):
        if not url or not key:
            raise SystemExit(
                "Faltam SUPABASE_URL e SUPABASE_SERVICE_KEY no ambiente.")
        self.base = url.rstrip("/") + "/rest/v1"
        self.h = {
            "apikey": key,
            "Authorization": f"Bearer {key}",
            "Content-Type": "application/json",
            "Accept-Profile": SCHEMA,     # leitura no schema gestao
            "Content-Profile": SCHEMA,    # escrita no schema gestao
        }

    def get(self, tabela, **params):
        r = requests.get(f"{self.base}/{tabela}", headers=self.h,
                         params=params, timeout=30)
        r.raise_for_status()
        return r.json()

    def patch(self, tabela, filtro, dados):
        r = requests.patch(f"{self.base}/{tabela}", headers=self.h,
                           params=filtro, json=dados, timeout=30)
        r.raise_for_status()
        return r.json() if r.text else []

    def insert(self, tabela, dados, upsert_on=None):
        h = dict(self.h)
        if upsert_on:
            # resolve conflito sem quebrar: e o caso do reprocessamento
            h["Prefer"] = f"resolution=merge-duplicates,return=representation"
        else:
            h["Prefer"] = "return=representation"
        params = {"on_conflict": upsert_on} if upsert_on else {}
        r = requests.post(f"{self.base}/{tabela}", headers=h,
                          params=params, json=dados, timeout=30)
        r.raise_for_status()
        return r.json() if r.text else []

    def evento(self, slug):
        e = self.get("eventos", slug=f"eq.{slug}", select="*", limit=1)
        if not e:
            raise SystemExit(f'Evento "{slug}" nao existe no banco.')
        return e[0]

    def rpc(self, funcao, args):
        r = requests.post(f"{self.base}/rpc/{funcao}", headers=self.h,
                          json=args, timeout=30)
        r.raise_for_status()
        return r.json() if r.text else None


# =====================================================================
# 1. SYMPLA  →  gestores + participantes
# =====================================================================

# A API do Sympla nao documenta os codigos de order_status (o proprio
# schema oficial da Sympla tipa o campo como string livre, sem enum).
# "A" foi confirmado batendo 100/100 pedidos aprovados de verdade, num
# evento real ja realizado (Experience 2026, id 3467585, em 02/09/2026)
# — os nomes por extenso ficam de fallback, caso outra conta/versao da
# API devolva isso diferente.
def _pedido_aprovado(order_status):
    return (order_status or "").strip().upper() in ("A", "APPROVED", "APROVADO", "COMPLETE")


# O link que fica em jantares.sympla_url segue o padrao
# https://www.sympla.com.br/slug-do-evento__<id>, com o id numerico do
# evento no Sympla depois do "__" — foi assim que os eventos "CIO
# Cerrado Experience 20XX" da propria conta vieram na API (checado em
# 02/09/2026: 1617982, 2017269, 2507590... todos com esse formato).
def _sympla_event_id_do_link(url):
    if not url:
        return None
    m = re.search(r"__(\d+)(?:[/?#].*)?$", url.strip())
    return m.group(1) if m else None


def sincronizar_sympla(supa, evento, producao):
    if not SYMPLA_TOKEN:
        log.warning("SYMPLA_TOKEN ausente: pulando sincronizacao.")
        return

    sympla_id = evento.get("sympla_event_id")
    if not sympla_id:
        log.warning("Evento sem sympla_event_id: cadastre no painel (aba Estrutura).")
        return

    log.info("Lendo participantes do Sympla (evento %s)...", sympla_id)

    participantes, pagina = [], 1
    while True:
        r = requests.get(
            f"https://api.sympla.com.br/public/v3/events/{sympla_id}/participants",
            headers={"s_token": SYMPLA_TOKEN},
            params={"page": pagina, "page_size": 200}, timeout=30)
        r.raise_for_status()
        corpo = r.json()
        lote = corpo.get("data", [])
        participantes.extend(lote)

        pg = corpo.get("pagination", {})
        if not pg.get("has_next") or not lote:
            break
        pagina += 1

    log.info("%d inscricao(oes) no Sympla.", len(participantes))

    novos = atualizados = ignorados = 0

    for p in participantes:
        # cancelado no Sympla nao entra nem atualiza
        if not _pedido_aprovado(p.get("order_status")):
            ignorados += 1
            continue

        email = (p.get("email") or "").strip().lower()
        nome  = (p.get("first_name", "") + " " + p.get("last_name", "")).strip()

        # Sem e-mail nao ha como deduplicar nem enviar contrato.
        if not email or not nome:
            ignorados += 1
            continue

        custom = {c.get("name"): c.get("value")
                  for c in (p.get("custom_form") or [])}

        dados_gestor = {
            "nome": nome,
            "email": email,
            "empresa": custom.get("Empresa") or p.get("company"),
            "cargo": custom.get("Cargo"),
            "telefone": p.get("phone"),
            "origem": "importacao",
        }
        dados_gestor = {k: v for k, v in dados_gestor.items() if v}

        if not producao:
            log.info("  [seguro] %s <%s>", nome, email)
            continue

        g = supa.insert("gestores", dados_gestor, upsert_on="email_norm")
        if not g:
            g = supa.get("gestores", email=f"eq.{email}", select="id", limit=1)
        gestor_id = g[0]["id"]

        ja = supa.get("participantes",
                      evento_id=f"eq.{evento['id']}",
                      gestor_id=f"eq.{gestor_id}",
                      select="id,status", limit=1)

        if ja:
            atualizados += 1
        else:
            supa.insert("participantes", {
                "evento_id": evento["id"],
                "gestor_id": gestor_id,
                "sympla_id": str(p.get("id")),
                "tipo_ingresso": p.get("ticket_name"),
                # aprovacao continua sendo decisao humana no painel
                "status": "pendente",
                "origem": "sympla",
            })
            novos += 1

    log.info("Sympla: %d novo(s), %d ja existiam, %d ignorado(s).",
             novos, atualizados, ignorados)


# ---------------------------------------------------------------------
# 1b. SYMPLA  →  convidados de jantar
#
# Jantar nao tem evento_id (e standalone por desenho — jantares.html
# ja documenta isso). Cada jantar pode ter o proprio evento no Sympla,
# guardado como link em jantares.sympla_url. Por isso essa etapa
# ignora o --evento da linha de comando e olha TODOS os jantares
# ativos (planejado/confirmado) que tem link configurado.
#
# Reusa jantar_importar_convidados_sympla via RPC — a mesma funcao SQL
# que a tela de admin ja chama quando alguem sobe a planilha na mao,
# so que aqui a "planilha" vem pronta da API. Sem duplicar a regra de
# negocio (achar/criar gestor, aprovado vs cancelado, etc.) em Python.
# ---------------------------------------------------------------------
DIACRITICOS = re.compile(r"[̀-ͯ]")


def _normaliza_rotulo(s):
    s = unicodedata.normalize("NFD", s or "")
    s = DIACRITICOS.sub("", s)
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _campo_sympla(custom_form, *alternativas):
    """Mesma tolerancia da tela (achaSympla em jantares.html): casa por
    prefixo, sem acento/maiuscula — o rotulo do formulario custom as
    vezes vem com acento quebrado na API."""
    campos = [(_normaliza_rotulo(c.get("name", "")), c.get("value", "")) for c in custom_form]
    for alvo in alternativas:
        alvo_norm = _normaliza_rotulo(alvo)
        for chave, valor in campos:
            if chave.startswith(alvo_norm) and str(valor).strip():
                return str(valor).strip()
    return ""


def sincronizar_jantares_sympla(supa, evento, producao):
    if not SYMPLA_TOKEN:
        log.warning("SYMPLA_TOKEN ausente: pulando sincronizacao de jantares.")
        return

    jantares = supa.get("jantares", select="id,patrocinador_nome,sympla_url",
                        status="in.(planejado,confirmado)")
    alvo = [(j, _sympla_event_id_do_link(j.get("sympla_url"))) for j in jantares]
    alvo = [(j, sid) for j, sid in alvo if sid]

    if not alvo:
        log.info("Nenhum jantar ativo com link do Sympla configurado.")
        return

    log.info("%d jantar(es) com link do Sympla.", len(alvo))

    for jantar, sympla_id in alvo:
        log.info("  %s (Sympla #%s)...", jantar["patrocinador_nome"], sympla_id)

        participantes, pagina = [], 1
        while True:
            r = requests.get(
                f"https://api.sympla.com.br/public/v3/events/{sympla_id}/participants",
                headers={"s_token": SYMPLA_TOKEN},
                params={"page": pagina, "page_size": 200}, timeout=30)
            if r.status_code == 404:
                log.warning("    evento %s nao encontrado no Sympla — confira o link", sympla_id)
                break
            r.raise_for_status()
            corpo = r.json()
            lote = corpo.get("data", [])
            participantes.extend(lote)
            pg = corpo.get("pagination", {})
            if not pg.get("has_next") or not lote:
                break
            pagina += 1

        if not participantes:
            log.info("    0 inscricao(oes).")
            continue

        linhas = []
        for p in participantes:
            custom = p.get("custom_form") or []
            nome = _campo_sympla(custom, "Nome Cracha", "Nome") or \
                   (p.get("first_name", "") + " " + p.get("last_name", "")).strip()
            linhas.append({
                "nome": nome,
                "email": (p.get("email") or "").strip().lower(),
                "email_corporativo": _campo_sympla(custom, "E-MAIL CORPORATIVO", "Email Corporativo"),
                "empresa": _campo_sympla(custom, "Empresa"),
                "cargo": _campo_sympla(custom, "Cargo"),
                "telefone": _campo_sympla(custom, "Telefone Celular", "Telefone"),
                "cpf": _campo_sympla(custom, "CPF"),
                "cnpj": _campo_sympla(custom, "CNPJ"),
                "sympla_id": str(p.get("id")),
                "estado_pagamento": "aprovado" if _pedido_aprovado(p.get("order_status")) else "cancelado",
            })

        if not producao:
            aprov = sum(1 for l in linhas if l["estado_pagamento"] == "aprovado")
            log.info("    [seguro] %d linha(s), %d aprovada(s)", len(linhas), aprov)
            continue

        try:
            r = supa.rpc("jantar_importar_convidados_sympla",
                        {"p_jantar_id": jantar["id"], "p_linhas": linhas})
            log.info("    %d novo(s), %d atualizado(s), %d recusado(s), "
                     "%d gestor(es) novo(s), %d erro(s) · %d/%d vaga(s)",
                     r.get("criados", 0), r.get("atualizados", 0), r.get("recusados", 0),
                     r.get("gestores_novos", 0), r.get("erros", 0),
                     r.get("ocupados_agora", 0), r.get("capacidade", 0))
        except Exception as e:
            log.error("    falhou: %s", e)


# =====================================================================
# 2. AUTENTIQUE  →  envio e status
# =====================================================================
AUTENTIQUE_API = "https://api.autentique.com.br/v2/graphql"


def autentique(query, variaveis=None):
    r = requests.post(
        AUTENTIQUE_API,
        headers={"Authorization": f"Bearer {AUTENTIQUE_TOKEN}"},
        json={"query": query, "variables": variaveis or {}}, timeout=40)
    r.raise_for_status()
    corpo = r.json()
    if corpo.get("errors"):
        raise RuntimeError(corpo["errors"][0].get("message", "erro Autentique"))
    return corpo.get("data", {})


# ---------------------------------------------------------------------
# Geracao do PDF a partir do .docx — mesmo mecanismo do pipeline antigo
# (cerrado_contratos.py): preenche nome/CPF por substituicao de
# paragrafo, converte via LibreOffice (soffice) ou, se nao tiver,
# docx2pdf (precisa do Word instalado — so funciona no Windows).
# ---------------------------------------------------------------------
MUTATION_CRIAR_DOCUMENTO = """
mutation CreateDocumentMutation(
  $document: DocumentInput!,
  $signers: [SignerInput!]!,
  $file: Upload!,
  $sandbox: Boolean
) {
  createDocument(document: $document, signers: $signers, file: $file, sandbox: $sandbox) {
    id
    name
    signatures { public_id email link { short_link } }
  }
}
"""


def _preencher_docx(template, destino, nome, cpf=""):
    from docx import Document
    doc = Document(template)
    ja_nome = ja_cpf = False
    for p in doc.paragraphs:
        texto = p.text.strip()
        if texto.startswith("Prezado(a) participante"):
            _substituir_paragrafo(p, f"Prezado(a) {nome},")
        elif texto.startswith("Nome:") and "_" in texto and not ja_nome:
            _substituir_paragrafo(p, f"Nome: {nome}")
            ja_nome = True
        elif texto.startswith("CPF:") and "_" in texto and not ja_cpf:
            if cpf:
                _substituir_paragrafo(p, f"CPF: {cpf}")
            ja_cpf = True
    doc.save(destino)


def _substituir_paragrafo(paragrafo, novo_texto):
    if paragrafo.runs:
        paragrafo.runs[0].text = novo_texto
        for run in paragrafo.runs[1:]:
            run.text = ""
    else:
        paragrafo.text = novo_texto


def gerar_contrato_pdf(nome, cpf=""):
    """Preenche o .docx com nome/CPF e converte pra PDF, em pasta
    temporaria. Levanta excecao clara se o template ou o conversor
    (LibreOffice/docx2pdf) nao estiverem disponiveis — quem chama
    decide se isso pula so este contrato ou para tudo."""
    import shutil, subprocess, tempfile, unicodedata

    if not CONTRATO_TEMPLATE_PATH or not os.path.exists(CONTRATO_TEMPLATE_PATH):
        raise RuntimeError(
            f"CONTRATO_TEMPLATE_PATH nao aponta pra um .docx que existe: "
            f"{CONTRATO_TEMPLATE_PATH!r}")

    nome_ascii = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode()
    nome_ascii = nome_ascii.replace(" ", "_") or "contrato"

    pasta_tmp = tempfile.mkdtemp(prefix="cerrado_")
    try:
        docx_tmp = os.path.join(pasta_tmp, f"{nome_ascii}.docx")
        pdf_tmp = os.path.join(pasta_tmp, f"{nome_ascii}.pdf")
        _preencher_docx(CONTRATO_TEMPLATE_PATH, docx_tmp, nome, cpf)

        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf",
                 "--outdir", pasta_tmp, docx_tmp],
                check=True, capture_output=True, timeout=120)
        else:
            from docx2pdf import convert
            convert(docx_tmp, pdf_tmp)

        if not os.path.exists(pdf_tmp):
            raise RuntimeError("PDF nao foi gerado na conversao")

        # move pra fora da pasta temporaria antes dela ser apagada
        destino = os.path.join(tempfile.gettempdir(), f"{nome_ascii}.pdf")
        shutil.move(pdf_tmp, destino)
        return destino
    finally:
        shutil.rmtree(pasta_tmp, ignore_errors=True)


def criar_documento_autentique(nome_doc, caminho_pdf, email, sandbox=True):
    """Sobe o PDF preenchido pro Autentique (createDocument, nao
    createDocumentFromTemplate — nao existe template cadastrado la).
    Um signatario so (o gestor) — o modelo de dados do gestao nao
    rastreia testemunha/co-signer, diferente do pipeline antigo que
    tinha Gerardo e Kelson fixos. Sem posicionamento automatico do
    carimbo de assinatura: depende do texto exato do template (o
    pipeline antigo procurava "Participante / CIO" etc. no PDF), e
    nao ha template de 2027 ainda pra saber se essas marcas existem —
    o signatario posiciona a propria assinatura na tela do Autentique."""
    documento = {"name": nome_doc}
    if AUTENTIQUE_DEADLINE:
        documento["deadline_at"] = AUTENTIQUE_DEADLINE

    operations = json.dumps({
        "query": MUTATION_CRIAR_DOCUMENTO,
        "variables": {
            "document": documento,
            "signers": [{"email": email, "action": "SIGN"}],
            "file": None,
            "sandbox": sandbox,
        },
    })
    file_map = json.dumps({"file": ["variables.file"]})

    with open(caminho_pdf, "rb") as f:
        r = requests.post(
            AUTENTIQUE_API,
            headers={"Authorization": f"Bearer {AUTENTIQUE_TOKEN}"},
            data={"operations": operations, "map": file_map},
            files={"file": (os.path.basename(caminho_pdf), f, "application/pdf")},
            timeout=60)
    r.raise_for_status()
    resp = r.json()
    if resp.get("errors"):
        raise RuntimeError(f"Erro Autentique: {resp['errors']}")
    return resp["data"]["createDocument"]


def enviar_contratos(supa, evento, producao):
    if not AUTENTIQUE_TOKEN:
        log.warning("AUTENTIQUE_TOKEN ausente: pulando envio.")
        return
    if not CONTRATO_TEMPLATE_PATH:
        log.warning("CONTRATO_TEMPLATE_PATH ausente: pulando envio.")
        return

    # aprovados que ainda nao receberam contrato
    pend = supa.get(
        "contratos",
        select="id,participante_id,status,participantes!inner(evento_id,status,gestores!inner(nome,email,cpf))",
        status="eq.nao_enviado")

    alvo = [c for c in pend
            if c["participantes"]["evento_id"] == evento["id"]
            and c["participantes"]["status"] == "aprovado"]

    log.info("%d contrato(s) para enviar.", len(alvo))

    for c in alvo:
        g = c["participantes"]["gestores"]
        if not producao:
            log.info("  [seguro] contrato para %s <%s>", g["nome"], g["email"])
            continue

        pdf = None
        try:
            pdf = gerar_contrato_pdf(g["nome"], g.get("cpf") or "")
            doc = criar_documento_autentique(
                f"Contrato · {g['nome']}", pdf, g["email"], sandbox=not producao)

            supa.patch("contratos", {"id": f"eq.{c['id']}"}, {
                "autentique_id": doc["id"],
                "autentique_url": f"https://app.autentique.com.br/documentos/{doc['id']}",
                "status": "enviado",
                "enviado_em": agora(),
            })
            enfileirar(supa, evento, g["email"], "contrato_enviado",
                       "Seu contrato do CIO Cerrado Experience")
            log.info("  enviado: %s", g["email"])

        except Exception as e:
            # um contrato que falha nao pode parar a fila inteira
            log.error("  falhou para %s: %s", g["email"], e)
        finally:
            if pdf and os.path.exists(pdf):
                try:
                    os.remove(pdf)
                except OSError:
                    pass


def ler_status(supa, evento, producao):
    if not AUTENTIQUE_TOKEN:
        log.warning("AUTENTIQUE_TOKEN ausente: pulando leitura de status.")
        return

    abertos = supa.get("contratos", select="id,autentique_id,participante_id",
                       status="eq.enviado")
    abertos = [c for c in abertos if c.get("autentique_id")]
    log.info("Conferindo %d contrato(s) em aberto...", len(abertos))

    assinados = 0
    for c in abertos:
        try:
            d = autentique("""
              query($id: UUID!) {
                document(id: $id) { id signatures { signed { created_at } } }
              }""", {"id": c["autentique_id"]})

            assinaturas = d["document"]["signatures"]
            if any(s.get("signed") for s in assinaturas):
                if producao:
                    supa.patch("contratos", {"id": f"eq.{c['id']}"},
                               {"status": "assinado", "assinado_em": agora()})
                    avisar_participante(supa, evento, c["participante_id"],
                                        "contrato_assinado",
                                        "Contrato assinado — complete sua hospedagem")
                assinados += 1
        except Exception as e:
            log.error("  erro no documento %s: %s", c["autentique_id"], e)

    log.info("%d assinatura(s) nova(s).", assinados)


# =====================================================================
# 3. LEMBRETES
# =====================================================================
def lembretes(supa, evento, producao):
    """Cobra quem recebeu o contrato e nao assinou."""
    hoje = datetime.now(timezone.utc)
    corte_primeiro = hoje - timedelta(days=LEMBRETE_APOS_DIAS)
    corte_repeticao = hoje - timedelta(days=LEMBRETE_INTERVALO)

    pend = supa.get(
        "contratos",
        select="id,lembretes_enviados,ultimo_lembrete_em,enviado_em,"
               "participantes!inner(evento_id,gestores!inner(nome,email))",
        status="eq.enviado")

    prazo = evento.get("prazo_contrato")
    enviados = 0

    for c in pend:
        if c["participantes"]["evento_id"] != evento["id"]:
            continue
        if (c.get("lembretes_enviados") or 0) >= LEMBRETE_MAXIMO:
            continue
        if not c.get("enviado_em"):
            continue

        enviado = iso(c["enviado_em"])
        if enviado > corte_primeiro:
            continue                       # ainda e cedo para cobrar

        ultimo = iso(c["ultimo_lembrete_em"]) if c.get("ultimo_lembrete_em") else None
        if ultimo and ultimo > corte_repeticao:
            continue                       # cobrado ha pouco

        g = c["participantes"]["gestores"]
        assunto = "Lembrete: seu contrato ainda não foi assinado"
        if prazo:
            assunto += f" (prazo {br(prazo)})"

        if not producao:
            log.info("  [seguro] lembrete para %s", g["email"])
            continue

        enfileirar(supa, evento, g["email"], "contrato_lembrete", assunto)
        supa.patch("contratos", {"id": f"eq.{c['id']}"}, {
            "lembretes_enviados": (c.get("lembretes_enviados") or 0) + 1,
            "ultimo_lembrete_em": agora(),
        })
        enviados += 1

    log.info("%d lembrete(s) enfileirado(s).", enviados)


# =====================================================================
# 4. FILA DE E-MAILS  →  Resend
# =====================================================================
CORPO = {
    "contrato_enviado":
        "Olá, {nome}.\n\n"
        "Seu contrato do CIO Cerrado Experience já está disponível para assinatura. "
        "Assim que assinar, o preenchimento dos dados de hospedagem é liberado.\n\n"
        "{link}\n\nEquipe CIO Cerrado",
    "contrato_lembrete":
        "Olá, {nome}.\n\n"
        "Seu contrato ainda está aguardando assinatura. Sem ela não conseguimos "
        "confirmar sua hospedagem.\n\n{link}\n\nEquipe CIO Cerrado",
    "contrato_assinado":
        "Olá, {nome}.\n\n"
        "Contrato assinado, obrigado. Agora complete seus dados de hospedagem "
        "— quem vai com você e se vai usar o transfer.\n\n{link}\n\nEquipe CIO Cerrado",
    "inscricao_aprovada":
        "Olá, {nome}.\n\n"
        "Sua inscrição no CIO Cerrado Experience foi aprovada. "
        "Em instantes você recebe o contrato para assinatura.\n\nEquipe CIO Cerrado",
    "autocadastro_recebido":
        "Olá, {nome}.\n\n"
        "Recebemos seu cadastro. Nossa equipe avalia e responde por aqui.\n\n"
        "Equipe CIO Cerrado",
    "rooming_ok":
        "Olá, {nome}.\n\n"
        "Seus dados de hospedagem foram confirmados. "
        "Se precisar mudar algo, é só voltar ao link.\n\n{link}\n\nEquipe CIO Cerrado",
}

PADRAO = "Olá, {nome}.\n\n{assunto}\n\nEquipe CIO Cerrado"


def processar_emails(supa, evento, producao, limite=200):
    if not RESEND_KEY:
        log.warning("RESEND_API_KEY ausente: pulando envio de e-mail.")
        return

    fila = supa.get("notificacoes", select="*",
                    status="eq.enfileirada",
                    evento_id=f"eq.{evento['id']}",
                    order="created_at.asc", limit=limite)

    log.info("%d e-mail(s) na fila.", len(fila))
    ok = falhou = 0

    for n in fila:
        nome = (n.get("destinatario") or "").split("@")[0]
        modelo = CORPO.get(n["tipo"], PADRAO)
        texto = modelo.format(nome=nome,
                              assunto=n.get("assunto") or "",
                              link=link_do_tipo(n["tipo"]))

        if not producao:
            log.info("  [seguro] %s -> %s", n["tipo"], n["destinatario"])
            continue

        try:
            r = requests.post(
                "https://api.resend.com/emails",
                headers={"Authorization": f"Bearer {RESEND_KEY}"},
                json={
                    "from": f"CIO Cerrado <{REMETENTE}>",
                    "to": [n["destinatario"]],
                    "subject": n.get("assunto") or "CIO Cerrado Experience",
                    "text": texto,
                }, timeout=30)
            r.raise_for_status()

            supa.patch("notificacoes", {"id": f"eq.{n['id']}"},
                       {"status": "enviada", "enviada_em": agora()})
            ok += 1

        except Exception as e:
            # marca o erro na propria linha: assim da para ver no painel
            # por que aquele e-mail nunca chegou
            supa.patch("notificacoes", {"id": f"eq.{n['id']}"},
                       {"status": "erro", "erro": str(e)[:500]})
            falhou += 1
            log.error("  falhou %s: %s", n["destinatario"], e)

    log.info("E-mails: %d enviado(s), %d com erro.", ok, falhou)


def link_do_tipo(tipo):
    base = os.environ.get("CERRADO_SITE", "https://ciocerrado.netlify.app")
    if tipo.startswith("contrato"):
        return f"{base}/rooming.html?evento={EVENTO_SLUG}"
    return f"{base}/rooming.html?evento={EVENTO_SLUG}"


# =====================================================================
# APOIO
# =====================================================================
def agora():
    return datetime.now(timezone.utc).isoformat()


def iso(v):
    """Aceita os dois formatos que o PostgREST devolve."""
    if not v:
        return None
    v = v.replace("Z", "+00:00")
    d = datetime.fromisoformat(v)
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


def br(d):
    return "/".join(reversed(str(d)[:10].split("-")))


def enfileirar(supa, evento, email, tipo, assunto):
    supa.insert("notificacoes", {
        "evento_id": evento["id"], "destinatario": email,
        "tipo": tipo, "assunto": assunto, "status": "enfileirada"})


def avisar_participante(supa, evento, participante_id, tipo, assunto):
    p = supa.get("participantes", select="gestores!inner(email)",
                 id=f"eq.{participante_id}", limit=1)
    if p:
        enfileirar(supa, evento, p[0]["gestores"]["email"], tipo, assunto)


# =====================================================================
# MAIN
# =====================================================================
def main():
    ap = argparse.ArgumentParser(description="Integracoes do CIO Cerrado")
    ap.add_argument("--tudo",      action="store_true")
    ap.add_argument("--sympla",    action="store_true")
    ap.add_argument("--jantares",  action="store_true",
                    help="convidados de jantar, pelo link do Sympla de cada jantar")
    ap.add_argument("--contratos", action="store_true")
    ap.add_argument("--status",    action="store_true")
    ap.add_argument("--lembretes", action="store_true")
    ap.add_argument("--emails",    action="store_true")
    ap.add_argument("--producao",  action="store_true",
                    help="envia de verdade (padrao e modo seguro)")
    ap.add_argument("--evento", default=EVENTO_SLUG)
    args = ap.parse_args()

    if not any([args.tudo, args.sympla, args.jantares, args.contratos,
                args.status, args.lembretes, args.emails]):
        ap.print_help()
        return 1

    if not args.producao:
        log.warning("MODO SEGURO: nada sera enviado. Use --producao para valer.")

    supa = Supa(SUPABASE_URL, SUPABASE_SERVICE)
    evento = supa.evento(args.evento)
    log.info("Evento: %s (%s)", evento["nome"], evento["slug"])

    # A ordem importa: sincroniza, envia, confere, cobra e so entao
    # despacha a fila — assim tudo que foi gerado agora ja sai junto.
    passos = [
        (args.tudo or args.sympla,    "Sympla",    sincronizar_sympla),
        (args.tudo or args.jantares,  "Jantares",  sincronizar_jantares_sympla),
        (args.tudo or args.contratos, "Contratos", enviar_contratos),
        (args.tudo or args.status,    "Status",    ler_status),
        (args.tudo or args.lembretes, "Lembretes", lembretes),
        (args.tudo or args.emails,    "E-mails",   processar_emails),
    ]

    for ativo, nome, funcao in passos:
        if not ativo:
            continue
        log.info("--- %s ---", nome)
        try:
            funcao(supa, evento, args.producao)
        except Exception as e:
            # um passo que quebra nao pode impedir os seguintes
            log.error("%s falhou: %s", nome, e)

    log.info("Fim.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
